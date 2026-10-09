#!/usr/bin/env python3
"""Fill free slides 12–14 — EXPERT evaluation copy (not team stage runbook).

Audience: experts reviewing the solution by 11.10 (pptx + demo + repo).
Does NOT touch locked slides 7–11.
Stage timing / who-clicks → docs/RUNBOOK_PITCH.md only.
After every edit: bash docs/presentation/run_round_outside.sh
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.util import Pt

ROOT = Path(__file__).resolve().parent
PPTX = ROOT / "FSP_NIKITA_DRAFT.pptx"
ASSETS = ROOT / "assets"

TITLE_PT = 22.0
BODY_PT = 13.0

INK = RGBColor(0x2A, 0x27, 0x33)
TITLE_ACCENT = RGBColor(0x6B, 0x2D, 0x8B)
TITLE_ON_DARK = RGBColor(0xFF, 0xFF, 0xFF)

SLIDE_TARGETS = {
    12: {
        "title_name": "Заголовок 27",
        "body_name": "Объект 28",
        "title_rgb": TITLE_ACCENT,
        "picture": None,
        "clear_names": (),
        "title": "Ядро решения: обратный найм с доказательствами",
        "fill_height": True,
        "body": (
            "Закрываем механику ТЗ, а не витрину вакансий: работодатель сам находит "
            "подтверждённую категорию и выходит с оффером и прозрачной вилкой ₽.\n"
            "Цепочка ценности: тест навыка → грейд на сервере (не из резюме) → "
            "подбор с честными статусами → приглашение с вилкой → контакты только "
            "после согласия и только этой компании.\n"
            "Честность в выдаче: подтверждено / нет / не проверено. "
            "«Не проверено» не выдаём за полный матч — меньше ложных матчей.\n"
            "Без боевого ФСП путь жив (сценарий без привязки + демо-достижения). "
            "Собственная платформа, не парсер внешних досок.\n"
            "Эксперт проверяет сам: живой HTTPS-стенд + открытый код + пары аккаунтов "
            "в материалах сдачи — один сквозной проход без сопровождения команды.\n"
            "Проверить сейчас: https://85-137-26-131.sslip.io/"
        ),
    },
    13: {
        "title_name": "Заголовок 1",
        "body_name": "Объект 2",
        "title_rgb": TITLE_ON_DARK,
        "picture": None,
        "clear_names": (),
        "as_table": True,  # 4 aligned rows — not freeform two columns
        "title": "Доказательства на стенде",
        "body": "",  # filled by table builder
    },
    14: {
        "title_name": "Заголовок 2",
        "body_name": "Текст 6",
        "title_rgb": TITLE_ACCENT,
        "picture": "slide14_demo_cover.png",
        "clear_names": (),
        "title": "Границы MVP: ядро сдано, оболочка — дальше",
        "fill_height": True,
        "body": (
            "Уже для экспертизы:\n"
            "два кабинета · тест→грейд · подбор с объяснением · приглашение с вилкой ₽ · "
            "адресные контакты · сценарий без ФСП и с демо-ФСП · живой стенд · открытый код.\n"
            "Сознательно снаружи ядра (граница пилота, не «дыра»):\n"
            "боевой реестр ФСП / Keycloak — мок по ТЗ; боевой SMTP не поднимали; "
            "рыночная сертификация метрик подбора — впереди.\n"
            "Справа на кадре — кабинет работодателя: фильтры, карточки, доказательства навыков.\n"
            "После отбора: боевой ФСП, почта, проверка на реальных работодателях — "
            "на уже работающем контуре согласия и грейда.\n"
            "Стенд сейчас: https://85-137-26-131.sslip.io/"
        ),
    },
}


def _shape_by_name(slide, name: str):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    raise RuntimeError(f"shape not found: {name!r}")


def _set_text(
    shape,
    text: str,
    *,
    size_pt: float,
    rgb: RGBColor,
    bold: bool | None = None,
    space_after_pt: float | None = None,
    fill_height: bool = False,
) -> None:
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size_pt)
        run.font.color.rgb = rgb
        if bold is not None:
            run.font.bold = bold
        p.space_before = Pt(0)
        p.space_after = Pt(0)
    if fill_height and len(lines) > 1:
        # Spread lines through the tall box (top-align alone leaves a white foot)
        box_h_pt = (shape.height / 914400.0) * 72.0
        line_h_pt = size_pt * 1.25
        # crude wrap: long lines consume extra line-heights
        used = 0.0
        w_in = max(shape.width / 914400.0, 1.0)
        cpl = max(8, int(w_in / (0.52 * size_pt / 72.0)))
        for line in lines:
            if not line.strip():
                used += line_h_pt * 0.4
            else:
                used += line_h_pt * max(1, (len(line) + cpl - 1) // cpl)
        # Use nearly full box; 0.88–0.92 left a visible white foot in PPT
        remain = max(0.0, box_h_pt * 0.97 - used)
        gap = remain / (len(lines) - 1)
        for p in tf.paragraphs[:-1]:
            p.space_after = Pt(max(2.0, gap))
        if tf.paragraphs:
            tf.paragraphs[-1].space_after = Pt(0)
    elif space_after_pt is not None:
        for p in tf.paragraphs[:-1]:
            p.space_after = Pt(space_after_pt)


def _clear_text(shape) -> None:
    if not shape.has_text_frame:
        return
    tf = shape.text_frame
    tf.clear()
    if tf.paragraphs and not tf.paragraphs[0].runs:
        tf.paragraphs[0].add_run().text = ""


def _fill_picture_slot(slide, path: Path) -> None:
    """Same pattern as insert_assets.py: drop placeholder, add_picture at slot."""
    if not path.is_file():
        raise FileNotFoundError(path)
    slot = None
    for sh in list(slide.shapes):
        try:
            if sh.is_placeholder and sh.placeholder_format.type == PP_PLACEHOLDER.PICTURE:
                slot = (sh.left, sh.top, sh.width, sh.height)
                sh._element.getparent().remove(sh._element)
        except Exception:
            continue
    for sh in list(slide.shapes):
        try:
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE and sh.left is not None and int(sh.left) > 5_000_000:
                sh._element.getparent().remove(sh._element)
        except Exception:
            continue
    if slot is None:
        slot = (6_096_000, 0, 6_096_000, 6_858_000)
    left, top, width, height = slot
    slide.shapes.add_picture(str(path), left, top, width=width, height=height)


def _remove_named(slide, name: str) -> None:
    for sh in list(slide.shapes):
        if sh.name == name:
            sh._element.getparent().remove(sh._element)


def _ensure_slide13_sidebar(slide, body_sh, sidebar: dict) -> None:
    """Split wide body into two script columns; stretch height to white-panel bottom."""
    _remove_named(slide, "Шпаргалка 13")
    _remove_named(slide, "Шпаргалка тело 13")
    gap = int(0.28 * 914400)
    full_w = int(11.80 * 914400)
    # Template white card (~Скругленный прямоугольник 3) bottom ≈ 6.95";
    # stock Объект 2 stopped ~1.15" early → permanent bottom desert.
    panel_bottom = int(6.90 * 914400)
    body_sh.height = max(body_sh.height, panel_bottom - body_sh.top)
    col_w = (full_w - gap) // 2
    body_sh.width = col_w
    rail_left = body_sh.left + col_w + gap
    title_h = int(0.36 * 914400)
    title_box = slide.shapes.add_textbox(rail_left, body_sh.top, col_w, title_h)
    title_box.name = "Шпаргалка 13"
    _set_text(
        title_box,
        sidebar["title"],
        size_pt=16.0,
        rgb=TITLE_ACCENT,
        bold=True,
        space_after_pt=4.0,
    )
    body_box = slide.shapes.add_textbox(
        rail_left,
        body_sh.top + title_h + int(0.04 * 914400),
        col_w,
        body_sh.height - title_h - int(0.04 * 914400),
    )
    body_box.name = "Шпаргалка тело 13"
    _set_text(
        body_box,
        sidebar["body"],
        size_pt=15.0,
        rgb=INK,
        bold=False,
        fill_height=True,
    )


def _fill_slide13_table(slide) -> None:
    """Real 5×2 table (header + 4 proof rows) — freestyle columns looked crooked."""
    # Remove prior freeform sidebar / old table from re-runs
    for name in ("Шпаргалка 13", "Шпаргалка тело 13", "Таблица 13"):
        for sh in list(slide.shapes):
            if getattr(sh, "name", "") == name:
                sh._element.getparent().remove(sh._element)
    for sh in list(slide.shapes):
        try:
            if sh.has_table and getattr(sh, "name", "").startswith("Таблица"):
                sh._element.getparent().remove(sh._element)
        except Exception:
            pass

    # Remove freeform body placeholder entirely — a shrunk box left a dashed remnant on capture
    try:
        body_sh = _shape_by_name(slide, "Объект 2")
        body_sh._element.getparent().remove(body_sh._element)
    except RuntimeError:
        pass  # already removed on re-run

    rows = [
        ("Шаг", "Что видно / что доказывает"),
        (
            "1 · Профиль",
            "Кандидат проходит тест; грейд из теста, не из резюме.",
        ),
        (
            "2 · Подбор",
            "HR видит основания «почему в выдаче» и честные статусы.",
        ),
        (
            "3 · Вилка ₽",
            "Адресное приглашение с вилкой до переписки, без обязательной вакансии.",
        ),
        (
            "4 · Согласие",
            "Контакты у пригласившего HR; у другого работодателя — нет. "
            "Стенд: https://85-137-26-131.sslip.io/",
        ),
    ]
    left = int(0.55 * 914400)
    top = int(1.70 * 914400)
    width = int(12.20 * 914400)
    height = int(4.90 * 914400)
    table_shape = slide.shapes.add_table(len(rows), 2, left, top, width, height)
    table_shape.name = "Таблица 13"
    table = table_shape.table
    table.columns[0].width = int(2.80 * 914400)
    table.columns[1].width = int(9.40 * 914400)
    from pptx.oxml.ns import qn

    for r, (c0, c1) in enumerate(rows):
        for c, text in enumerate((c0, c1)):
            cell = table.cell(r, c)
            cell.text = text
            try:
                cell.text_frame.margin_left = int(0.10 * 914400)
                cell.text_frame.margin_right = int(0.10 * 914400)
                cell.text_frame.margin_top = int(0.08 * 914400)
                cell.text_frame.margin_bottom = int(0.08 * 914400)
            except Exception:
                pass
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(15 if r == 0 else 13)
                    run.font.bold = r == 0 or c == 0
                    run.font.color.rgb = TITLE_ACCENT if (r == 0 or c == 0) else INK
    # White panel hugs the table; strip stroke (read as dashed frame on capture)
    try:
        panel = _shape_by_name(slide, "Скругленный прямоугольник 3")
        panel.top = int(1.50 * 914400)
        panel.height = int(5.30 * 914400)
        panel.left = int(0.36 * 914400)
        panel.width = int(12.62 * 914400)
        spPr = panel._element.spPr
        ln = spPr.find(qn("a:ln")) if spPr is not None else None
        if ln is not None:
            spPr.remove(ln)
    except RuntimeError:
        pass


def fill_slide(prs: Presentation, sn: int, payload: dict) -> None:
    slide = prs.slides[sn - 1]
    for name in payload.get("clear_names") or ():
        try:
            _clear_text(_shape_by_name(slide, name))
        except RuntimeError:
            pass
    title_sh = _shape_by_name(slide, payload["title_name"])
    # Titles on 12/14 were glued to the top edge — drop to ~1.35″
    if sn in (12, 14):
        title_sh.top = int(1.35 * 914400)
    if sn == 13:
        title_sh.top = int(0.85 * 914400)
    _set_text(
        title_sh,
        payload["title"],
        size_pt=TITLE_PT,
        rgb=payload["title_rgb"],
        bold=True,
    )
    if payload.get("as_table"):
        _fill_slide13_table(slide)
        print(f"filled slide {sn}: {payload['title']} (table)")
        return

    body_sh = _shape_by_name(slide, payload["body_name"])
    body_pt = float(payload.get("body_pt", BODY_PT))
    # Slide 12/14: hug body to copy so fill_height doesn't invent a white foot
    if sn in (12, 14):
        lines = [ln for ln in payload["body"].split("\n") if ln.strip()]
        w_in = max(body_sh.width / 914400.0, 1.0)
        cpl = max(8, int(w_in / (0.52 * body_pt / 72.0)))
        used_lines = 0
        for ln in lines:
            used_lines += max(1, (len(ln) + cpl - 1) // cpl)
        need_in = used_lines * (body_pt / 72.0) * 1.40 + 0.25
        lo, hi = (3.2, 4.6) if sn == 12 else (2.8, 4.0)
        body_sh.height = int(max(lo, min(hi, need_in)) * 914400)
        # Keep body below lowered title (absolute — re-runs must not drift)
        body_sh.top = title_sh.top + title_sh.height + int(0.28 * 914400)
    if sn == 14:
        try:
            panel = _shape_by_name(slide, "Скругленный прямоугольник 8")
            panel.top = int(0.45 * 914400)
            panel.height = (body_sh.top + body_sh.height + int(0.28 * 914400)) - panel.top
        except RuntimeError:
            pass
    _set_text(
        body_sh,
        payload["body"],
        size_pt=body_pt,
        rgb=INK,
        bold=False,
        fill_height=bool(payload.get("fill_height")),
    )
    pic = payload.get("picture")
    if pic:
        _fill_picture_slot(slide, ASSETS / pic)
    print(f"filled slide {sn}: {payload['title']}")


def ensure_slide14_cover() -> Path:
    """Portrait overview of employer search — keep filters + cards, drop sparse page foot."""
    out = ASSETS / "slide14_demo_cover.png"
    # Prefer pre-framed overview (full cabinet). Landscape left-strip of desktop_search
    # cuts off candidate cards and looks like an empty filter panel.
    overview = ASSETS / "slide14_demo_cover_overview.png"
    src_path = overview if overview.is_file() else ASSETS / "demo_desktop_search.png"
    if not src_path.is_file():
        raise FileNotFoundError(src_path)
    from PIL import Image

    slot_aspect = 6.667 / 7.5
    src = Image.open(src_path).convert("RGB")
    w, h = src.size
    # Drop bottom ~18% (empty list foot under cards); then cover-fit to slot aspect
    band_h = int(h * 0.82)
    cover = src.crop((0, 0, w, band_h))
    cw, ch = cover.size
    cur = cw / ch
    if cur > slot_aspect:
        nw = int(ch * slot_aspect)
        x0 = (cw - nw) // 2
        cover = cover.crop((x0, 0, x0 + nw, ch))
    elif cur < slot_aspect:
        nh = int(cw / slot_aspect)
        cover = cover.crop((0, 0, cw, min(ch, nh)))
    cover.resize((1600, int(1600 / slot_aspect)), Image.Resampling.LANCZOS).save(
        out, optimize=True
    )
    return out


def main() -> None:
    ensure_slide14_cover()
    prs = Presentation(str(PPTX))
    for sn, payload in SLIDE_TARGETS.items():
        fill_slide(prs, sn, payload)
    prs.save(str(PPTX))
    print("saved", PPTX)


if __name__ == "__main__":
    main()
