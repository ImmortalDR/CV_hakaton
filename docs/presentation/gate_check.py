#!/usr/bin/env python3
"""Machine checks for LCT presentation draft — design-principles oriented.

Canon: Open_Portfolio_AI/growth/ui-ux-verification-2026.md (SG-*).
Lessons: docs/presentation/LESSONS_LCT_2026-10-09.md

- CRAP: Contrast, Repetition, Alignment, Proximity (content-gap, not tall-box ignore)
- Billboard: cut copy, do not shrink past Geoscan ramp
- Density-ref: same template boxes need Geoscan-scale length (anti-desert)
- Overflow FILL_MAX calibrated so Geoscan reference would not false-FAIL
- After every edit: run_round.sh (this file + full_audit + PPT capture)

Does NOT replace PowerPoint eye check (SG-channel / SG-comp). Exit 0 = GATE PASS.
"""

from __future__ import annotations

import re
import sys
import zipfile
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parent
PPTX = ROOT / "FSP_NIKITA_DRAFT.pptx"
EMU = 914400.0

OWN_SLIDE_INDEXES = (7, 8, 9, 10, 11)  # 1-based

FORBIDDEN = (
    "НАЗВАНИЕ",
    "Имя Фамилия",
    "Образец текста",
    "дописать",
    "не используется",
    "неиспользуем",
    "В чем суть вашего решения",
    "В чём суть вашего решения",
    "Что делает ваше решение уникальным",
    "Капитан: ФИО",
)

BAD_LONG_TITLE = "КОМАНДА «Нанотехнологический альянс»"

# --- Type ramp = Geoscan on this template (11–22) ---
BODY_PT = 13.0
TITLE_PT = 22.0
ALLOWED_FILLED_PT = {11.0, 12.0, 13.0, 15.0, 16.0, 20.0, 22.0}
MIN_BODY_PT = 11.0
MIN_ROLE_PT = 11.0
MIN_PANEL_PT = 13.0
PANEL_AREA_IN2 = 18.0

# --- Billboard ---
MAX_BODY_LINES = {
    8: 5,
    9: 3,
    10: 4,
    11: 6,
}

# --- Proximity: measure CONTENT bottom of label (wrap×pt), not tall empty shape box ---
# Template nests body under tall label *shapes* (slide 8) — shape gap can be negative.
# Visual collision = body.top < label_content_bottom + MIN. That is what humans see.
PROX_CONTENT_GAP_MIN = 0.06  # inches of air between last label line and body top
PROX_CONTENT_GAP_MAX = 0.85
# Cyrillic bold labels wrap wider than body estimate
LABEL_CHAR_W = 0.62

# --- Overflow ---
# Calibrated on Geoscan slide2 same 0.97" boxes: their essence needs~1.45" by this
# estimator and still fits in PowerPoint. Old 1.06 false-failed the reference deck.
FILL_MAX = 1.55  # need/have height

# Bodies subject to billboard / overflow / type checks
BODY_KEYS = (
    "Тест навыков",
    "Грейд из теста",
    "Участники:",
    "Состав:",
    "Тимлид",
    "UI и сценарии",
    "Тест и подбор",
    "Собрались",
    "альянс» собрался",
    "Сначала навыки",
    "Нужно связать",
    "Вилка",
    "агрегатор",
    "Главный вызов",
    "React",
    "Стек.",
    "Путь.",
    "«Нанотехнологический",
)

# Label → body pairing for proximity (slide, label_prefix, body_key)
PROXIMITY_PAIRS = (
    (8, "Краткое описание", "Тест навыков"),
    (8, "Уникальность", "Грейд из теста"),
    (8, "О команде", "Участники:"),
    (10, "Краткая история", "альянс» собрался"),
    (10, "Почему вы", "Нужно связать"),
    (10, "С какими", "Главный вызов"),
    (11, "Техническая суть", "Стек."),
    (11, "Почему это выгодно", "Путь."),
)

# Cards: Geoscan centered trio (~2.52 … ~8.42), not edge-slammed
CARD_LEFT_MAX = 2.8
CARD_RIGHT_MIN = 8.0
CARD_LEFT_MIN = 2.2  # not stuck at far left margin

# SG-density-ref: min chars vs Geoscan on same template (anti-desert).
# Floors are ~85% of Geoscan measured lengths — allow edit, forbid sparse stubs.
DENSITY_FLOORS: dict[tuple[int, str], int] = {
    (8, "Тест навыков"): 240,
    (8, "Грейд из теста"): 220,
    (8, "Участники:"): 140,
    (10, "альянс» собрался"): 120,
    (10, "Нужно связать"): 120,
    (10, "Главный вызов"): 120,
    (11, "Стек."): 320,
    (11, "Путь."): 300,
}

LABEL_PREFIXES = (
    "Краткое описание",
    "Уникальность",
    "О команде",
    "Краткая история",
    "Почему вы",
    "С какими",
    "Техническая суть",
    "Маркетинговая суть",
)
# Note: «КОРОТКО О РЕШЕНИИ» is the pink title pill on slide 11 — not a section label.


def slide_texts_xml(z: zipfile.ZipFile, slide_path: str) -> list[str]:
    xml = z.read(slide_path).decode("utf-8", errors="replace")
    return [t for t in re.findall(r"<a:t[^>]*>([^<]*)</a:t>", xml) if t.strip()]


def layout_texts(z: zipfile.ZipFile) -> list[str]:
    out: list[str] = []
    for name in z.namelist():
        if name.startswith("ppt/slideLayouts/") and name.endswith(".xml"):
            out.extend(slide_texts_xml(z, name))
    return out


def presentation_order(z: zipfile.ZipFile) -> list[str]:
    xml = z.read("ppt/presentation.xml").decode("utf-8", errors="replace")
    rels = z.read("ppt/_rels/presentation.xml.rels").decode("utf-8", errors="replace")
    rid_to_target = dict(re.findall(r'Id="(rId\d+)"[^>]*Target="([^"]+)"', rels))
    rids = re.findall(r'<p:sldId[^>]*r:id="(rId\d+)"', xml) or re.findall(
        r'r:id="(rId\d+)"', xml
    )
    paths = []
    for rid in rids:
        tgt = rid_to_target.get(rid, "")
        if tgt:
            paths.append("ppt/" + tgt.lstrip("/"))
    return paths


def shape_area_in2(shape) -> float:
    try:
        return (shape.width * shape.height) / (EMU**2)
    except Exception:
        return 0.0


def shape_wh_in(shape) -> tuple[float, float]:
    return shape.width / EMU, shape.height / EMU


def first_run_pt(shape) -> float | None:
    if not shape.has_text_frame:
        return None
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            if run.text.strip() and run.font.size is not None:
                return float(run.font.size.pt)
    return None


def is_label(text: str) -> bool:
    t = text.strip()
    return any(t.startswith(p) for p in LABEL_PREFIXES) or len(t) < 12


def logical_lines(text: str) -> int:
    return len([ln for ln in text.split("\n") if ln.strip()])


def estimate_need_h_in(
    text: str, width_in: float, pt: float = 18.0, *, char_w_factor: float = 0.52
) -> float:
    avg_char_w = char_w_factor * pt / 72.0  # body 0.52; bold labels use LABEL_CHAR_W
    cpl = max(8, int(width_in / avg_char_w))
    lines = 0
    for para in text.split("\n"):
        if not para.strip():
            continue
        lines += max(1, (len(para) + cpl - 1) // cpl)
    return lines * (pt / 72.0) * 1.15


def label_content_bottom_emu(lab) -> int:
    """Estimated Y where label *glyphs* end (not the tall empty one-line box)."""
    text = lab.text_frame.text.strip()
    pt = first_run_pt(lab) or 16.0
    w_in = lab.width / EMU
    # Looser leading than body — matches PowerPoint multi-line section titles
    avg_char_w = LABEL_CHAR_W * pt / 72.0
    cpl = max(8, int(w_in / avg_char_w))
    lines = 0
    for para in text.split("\n"):
        if not para.strip():
            continue
        lines += max(1, (len(para) + cpl - 1) // cpl)
    need = lines * (pt / 72.0) * 1.35
    content = lab.top + max(need, pt / 72.0 * 1.35) * EMU
    shape_b = lab.top + lab.height
    if lines >= 2:
        return int(max(content, shape_b))
    return int(content)


def matches_body(text: str) -> bool:
    return any(k in text for k in BODY_KEYS)


def find_shape_with(slide, pred):
    for sh in slide.shapes:
        if getattr(sh, "has_text_frame", False) and pred(sh.text_frame.text.strip()):
            return sh
    return None


def check_slide8_is_product_ui(fails: list[str]) -> None:
    """Geoscan reference: branded team panel in left slot (dark purple + faces).

    Fail only if missing / tiny / empty monogram placeholder.
    """
    path = ROOT / "assets" / "slide8_visual.png"
    if not path.exists():
        fails.append("assets/slide8_visual.png missing")
        return
    try:
        from PIL import Image

        im = Image.open(path).convert("RGB")
        if im.size[0] < 800 or im.size[1] < 400:
            fails.append(f"SG-photo-slot slide8_visual too small {im.size}")
            return
        # Reject near-empty / single-letter monogram (very low color variance)
        small = im.resize((64, 32))
        pixels = list(small.getdata())
        mean = sum(p[0] + p[1] + p[2] for p in pixels) / (len(pixels) * 3)
        var = sum((p[0] + p[1] + p[2] - mean) ** 2 for p in pixels) / len(pixels)
        if var < 80:
            fails.append(
                f"SG-photo-slot slide8_visual looks blank/flat (var={var:.0f}) — need brand panel"
            )
    except Exception as e:
        fails.append(f"SG-photo-slot cannot inspect slide8_visual: {e}")


def check_empty_titles(prs: Presentation, fails: list[str]) -> None:
    """Banal: pink title pill empty OR dark text on pink (slides 8–10)."""
    from pptx.enum.shapes import PP_PLACEHOLDER

    for human_n in (8, 9, 10):
        slide = prs.slides[human_n - 1]
        found = False
        for sh in slide.shapes:
            try:
                if not sh.is_placeholder:
                    continue
                if sh.placeholder_format.type != PP_PLACEHOLDER.TITLE:
                    continue
                found = True
                text = sh.text_frame.text.strip() if sh.has_text_frame else ""
                if not text:
                    fails.append(
                        f"slide {human_n}: SG-title-empty — TITLE placeholder blank "
                        "(розовый заголовок без текста)"
                    )
                elif "НАЗВАНИЕ" in text:
                    fails.append(f"slide {human_n}: SG-title-empty still «НАЗВАНИЕ»")
                # Contrast depends on chrome: pink pill (9–10) → light; white field (8) → dark
                on_pink = False
                for other in slide.shapes:
                    oname = getattr(other, "name", "") or ""
                    if "Скругленный прямоугольник" not in oname:
                        continue
                    try:
                        if (
                            other.top < 1.2 * 914400
                            and other.left < 5.0 * 914400
                            and other.height < 1.0 * 914400
                        ):
                            on_pink = True
                            break
                    except Exception:
                        continue
                for para in sh.text_frame.paragraphs:
                    for run in para.runs:
                        if not run.text.strip():
                            continue
                        try:
                            c = run.font.color.rgb
                        except Exception:
                            c = None
                        if c is None:
                            fails.append(
                                f"slide {human_n}: SG-title-contrast — title has no explicit color"
                            )
                            continue
                        R, G, B = int(str(c)[0:2], 16), int(str(c)[2:4], 16), int(str(c)[4:6], 16)
                        lum = R + G + B
                        if on_pink and lum < 500:
                            fails.append(
                                f"slide {human_n}: SG-title-contrast — dark title {c} on pink pill"
                            )
                        if (not on_pink) and lum > 500:
                            fails.append(
                                f"slide {human_n}: SG-title-contrast — light title {c} on white "
                                "(нечитаемо)"
                            )
            except Exception:
                continue
        if not found:
            fails.append(f"slide {human_n}: SG-title-empty — no TITLE placeholder found")


def check_label_contrast(prs: Presentation, fails: list[str]) -> None:
    """Banal: section labels on theme color → white-on-white in export/projector."""
    for human_n in (8, 10, 11):
        slide = prs.slides[human_n - 1]
        for sh in slide.shapes:
            if not getattr(sh, "has_text_frame", False):
                continue
            text = sh.text_frame.text.strip()
            if not any(text.startswith(p) for p in LABEL_PREFIXES):
                continue
            ok = False
            for para in sh.text_frame.paragraphs:
                for run in para.runs:
                    if not run.text.strip():
                        continue
                    try:
                        c = run.font.color.rgb
                    except Exception:
                        c = None
                    if c is None:
                        fails.append(
                            f"slide {human_n}: SG-label-contrast — theme/no RGB on «{text[:40]}»"
                        )
                        continue
                    R, G, B = int(str(c)[0:2], 16), int(str(c)[2:4], 16), int(str(c)[4:6], 16)
                    if R + G + B > 520:
                        fails.append(
                            f"slide {human_n}: SG-label-contrast — too light {c} «{text[:40]}»"
                        )
                    else:
                        ok = True
                    pt = run.font.size.pt if run.font.size else None
                    if pt is not None and round(pt) not in (16, 20, 22):
                        fails.append(
                            f"slide {human_n}: SG-type-soup label {pt:.0f}pt «{text[:40]}»"
                        )
            if not ok and text:
                # already failed above if no rgb; silence if ok
                pass


def check_slide9_photo_square(prs: Presentation, fails: list[str]) -> None:
    """Circular portraits must sit in square frames — non-square = visible squash (SG-aspect)."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    slide = prs.slides[8]
    pics = [sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]
    if len(pics) < 3:
        fails.append(f"slide 9: SG-aspect expected ≥3 portraits, got {len(pics)}")
        return
    for sh in pics:
        w, h = sh.width / EMU, sh.height / EMU
        ratio = w / max(h, 1e-6)
        if abs(ratio - 1.0) > 0.04:
            fails.append(
                f"slide 9: SG-aspect portrait {w:.2f}×{h:.2f}\" ratio={ratio:.3f} "
                f"(must be square or circles squash)"
            )


def check_slide8_face_rings_no_overlap(fails: list[str]) -> None:
    """SG-overlap: circular faces on slide8_visual must not intersect.

    Only scans the face band (not name/role magenta text below), around three
    expected centers (design 300/700/1100 on 1400-wide panel).
    """
    path = ROOT / "assets" / "slide8_visual.png"
    if not path.exists():
        return
    try:
        from PIL import Image
    except ImportError:
        return
    im = Image.open(path).convert("RGB")
    w, h = im.size
    px = im.load()
    scale = w / 1400.0
    # Design: face_cy=270, diam=280 → ring only (exclude name/role text below)
    face_cy = int(270 * scale)
    diam = int(280 * scale)
    y0, y1 = face_cy - diam // 2 - 4, face_cy + diam // 2 + 4
    centers = [int(c * scale) for c in (300, 700, 1100)]
    bounds = []
    for cx in centers:
        xs = []
        for y in range(max(0, y0), min(h, y1), 1):
            for x in range(max(0, cx - diam // 2 - 8), min(w, cx + diam // 2 + 8), 1):
                r, g, b = px[x, y]
                if r > 200 and g < 90 and 40 < b < 140:
                    xs.append(x)
        if len(xs) < 30:
            fails.append(f"slide8_visual: SG-overlap missing face ring near x={cx}")
            return
        bounds.append((min(xs), max(xs)))
    for i in range(len(bounds) - 1):
        gap = bounds[i + 1][0] - bounds[i][1]
        if gap < 12:
            fails.append(
                f"slide8_visual: SG-overlap face rings gap={gap}px "
                f"(circles intersect)"
            )


def check_slide8_photo_clearance(prs: Presentation, fails: list[str]) -> None:
    """Slot must have a picture; Geoscan intentionally meets the divider (no fail on panel)."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    slide = prs.slides[7]
    pics = [
        sh
        for sh in slide.shapes
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE and sh.left is not None and sh.left < 1_000_000
    ]
    if not pics:
        fails.append("slide 8: SG-photo-slot — no picture in left slot")
        return
    pic = pics[0]
    h = pic.height / EMU
    if h < 2.8:
        fails.append(f"slide 8: SG-photo-slot height {h:.2f}\" too short (Geoscan ~3.25\")")


def check_banal_copy(prs: Presentation, fails: list[str]) -> None:
    """VT breaks, English Invite in narrative, etc."""
    import re

    # Own filled slides only (7 is organizer case title — leave unless we touch it)
    for human_n in (8, 9, 10, 11):
        slide = prs.slides[human_n - 1]
        blob_parts = []
        for sh in slide.shapes:
            if not getattr(sh, "has_text_frame", False):
                continue
            t = sh.text_frame.text
            if "\x0b" in t or "\v" in t:
                fails.append(
                    f"slide {human_n}: SG-vt — soft-break/U+000B in «{t.strip()[:40]}»"
                )
            blob_parts.append(t)
        blob = "\n".join(blob_parts)
        if human_n in (8, 10, 11) and re.search(r"\bInvite\b|\binvite\b", blob):
            fails.append(
                f"slide {human_n}: SG-copy-ru — «Invite/invite» in narrative; use «приглашение»"
            )


def check_layout_and_assets(fails: list[str]) -> None:
    with zipfile.ZipFile(PPTX) as z:
        order = presentation_order(z)
        if len(order) < max(OWN_SLIDE_INDEXES):
            fails.append(f"deck has {len(order)} slides, expected ≥{max(OWN_SLIDE_INDEXES)}")

        for human_n in OWN_SLIDE_INDEXES:
            if human_n - 1 >= len(order):
                continue
            path = order[human_n - 1]
            texts = slide_texts_xml(z, path)
            blob = "\n".join(texts)
            for bad in FORBIDDEN:
                if bad in blob:
                    if bad == "НАЗВАНИЕ" and "НАЗВАНИЕ" not in blob.replace("НТ АЛЬЯНС", ""):
                        continue
                    fails.append(f"slide {human_n}: forbidden «{bad}»")
            if BAD_LONG_TITLE in blob:
                fails.append(f"slide {human_n}: long title — use «НТ АЛЬЯНС»")
            # Repetition anti-pattern: manual bullets + template markers
            if re.search(r"(?m)^[•●▪]", blob) or "• " in blob:
                fails.append(
                    f"slide {human_n}: SG-repetition leading «•» — template already has markers"
                )

            if human_n in (8, 9):
                rels_path = path.replace("ppt/slides/", "ppt/slides/_rels/") + ".rels"
                if rels_path in z.namelist():
                    rels = z.read(rels_path).decode("utf-8", errors="replace")
                    media_n = len(re.findall(r"media/image", rels))
                    need = 1 if human_n == 8 else 3
                    if media_n < need:
                        fails.append(
                            f"slide {human_n}: only {media_n} media link(s), need ≥{need} photos"
                        )

        layouts = "\n".join(layout_texts(z))
        if "КОМАНДА «НАЗВАНИЕ»" in layouts or "НАЗВАНИЕ КОМАНДЫ" in layouts:
            fails.append("layouts: still contain «НАЗВАНИЕ» placeholder")
        if BAD_LONG_TITLE in layouts:
            fails.append("layouts: long team title — prefer «КОМАНДА «НТ АЛЬЯНС»»")
        if "НТ АЛЬЯНС" not in layouts:
            fails.append("layouts: missing «НТ АЛЬЯНС» team title")

    geo = ROOT / "assets" / "from_geoscan"
    if geo.exists():
        for name in ("portrait_danila.png", "portrait_nikita.png", "portrait_timofey.png"):
            p = ROOT / "assets" / name
            if not p.exists() or p.stat().st_size < 80_000:
                fails.append(f"assets/{name}: expected Geoscan photo (≥80KB), got missing/tiny")

    check_slide8_is_product_ui(fails)


def check_design_gates(prs: Presentation, fails: list[str]) -> None:
    for human_n in OWN_SLIDE_INDEXES:
        slide = prs.slides[human_n - 1]
        max_lines = MAX_BODY_LINES.get(human_n, 4)

        for sh in slide.shapes:
            if not getattr(sh, "has_text_frame", False):
                continue
            text = sh.text_frame.text.strip()
            if not text or is_label(text) or not matches_body(text):
                continue
            pt = first_run_pt(sh)
            area = shape_area_in2(sh)
            w_in, h_in = shape_wh_in(sh)
            nlines = logical_lines(text)

            # Type floor = Geoscan ramp on this template
            if pt is not None and matches_body(text):
                rounded = round(pt)
                if rounded < 11:
                    fails.append(
                        f"slide {human_n}: SG-contrast body {pt:.0f}pt < 11pt «{text[:40]}»"
                    )
                if rounded not in (11, 12, 13, 15, 16, 20, 22):
                    fails.append(
                        f"slide {human_n}: SG-type-soup {pt:.0f}pt off Geoscan ramp «{text[:40]}»"
                    )

            # Billboard: too many lines = wall of text
            if nlines > max_lines:
                fails.append(
                    f"slide {human_n}: SG-billboard {nlines} lines > {max_lines} "
                    f"«{text[:36]}…» — режь текст, не кегль"
                )

            # Overflow (cut, don't shrink past floor)
            if pt is not None and area >= 1.2:
                need = estimate_need_h_in(text, w_in, pt)
                if need > h_in * FILL_MAX:
                    fails.append(
                        f"slide {human_n}: SG-overflow need~{need:.2f}\" > have {h_in:.2f}\" "
                        f"«{text[:36]}…»"
                    )

            # Abandoned slot: huge field + tiny type + almost no lines
            # Geoscan often uses one dense paragraph (~13pt) in ~8–10 in² — that is OK.
            if area >= 14.0 and pt is not None and pt < 12 and nlines <= 1:
                fails.append(
                    f"slide {human_n}: SG-hierarchy desert {pt:.0f}pt / {nlines} lines "
                    f"in {area:.1f}in² «{text[:36]}…»"
                )

        # Proximity: label CONTENT → body (catches wrap into body; ignores tall empty boxes)
        slide_pairs = [(lp, bk) for sn, lp, bk in PROXIMITY_PAIRS if sn == human_n]
        resolved = []
        for lab_pref, body_key in slide_pairs:
            lab = find_shape_with(slide, lambda t, p=lab_pref: t.startswith(p))
            body = find_shape_with(slide, lambda t, k=body_key: k in t)
            if not lab or not body:
                fails.append(
                    f"slide {human_n}: SG-proximity missing pair «{lab_pref}» → «{body_key}»"
                )
                continue
            content_b = label_content_bottom_emu(lab)
            gap = (body.top - content_b) / EMU
            if gap < PROX_CONTENT_GAP_MIN - 1e-6:
                fails.append(
                    f"slide {human_n}: SG-overlap content-gap {gap:.2f}\" < {PROX_CONTENT_GAP_MIN}\" "
                    f"— тело наезжает на текст заголовка «{lab_pref}» "
                    f"(lab_content_b={content_b/EMU:.2f}\" body_t={body.top/EMU:.2f}\")"
                )
            if gap > PROX_CONTENT_GAP_MAX + 1e-6:
                fails.append(
                    f"slide {human_n}: SG-proximity content-gap {gap:.2f}\" > {PROX_CONTENT_GAP_MAX}\" "
                    f"— тело оторвано от заголовка «{lab_pref}»"
                )
            resolved.append((lab, body, lab_pref))

        # Body CONTENT must not invade a LOWER label in the SAME column (slide 10).
        # Skip cross-column pairs (slide 8 left «О команде» vs right essence; slide 11 two cols).
        for lab, body, lab_pref in resolved:
            pt = first_run_pt(body) or 13.0
            w_in = body.width / EMU
            need = estimate_need_h_in(body.text_frame.text.strip(), w_in, pt)
            body_content_b = body.top / EMU + need
            for other_lab, _other_body, other_pref in resolved:
                if other_lab is lab:
                    continue
                if other_lab.top <= body.top:
                    continue  # only labels below this body
                # same column = horizontal overlap of body box and other label box
                ox = min(body.left + body.width, other_lab.left + other_lab.width) - max(
                    body.left, other_lab.left
                )
                if ox < 0.5 * EMU:
                    continue
                gap2 = other_lab.top / EMU - body_content_b
                if gap2 < PROX_CONTENT_GAP_MIN - 1e-6:
                    fails.append(
                        f"slide {human_n}: SG-overlap body→next-label gap {gap2:.2f}\" "
                        f"— «{lab_pref}» body въезжает в «{other_pref}» "
                        f"(body_content_b={body_content_b:.2f}\" next_lab_t={other_lab.top/EMU:.2f}\")"
                    )

        # Card span (alignment across slide)
        if human_n == 9:
            cards = []
            for sh in slide.shapes:
                name = getattr(sh, "name", "") or ""
                if (
                    "Скругленный прямоугольник" in name
                    and sh.width > 2.0 * EMU
                    and sh.height > 4.0 * EMU
                ):
                    cards.append(sh.left / EMU)
            cards.sort()
            if len(cards) != 3:
                fails.append(f"slide 9: SG-card-span expected 3 cards, got {len(cards)}")
            else:
                if cards[0] > CARD_LEFT_MAX:
                    fails.append(
                        f"slide 9: SG-card-span leftmost {cards[0]:.2f}\" > {CARD_LEFT_MAX}\""
                    )
                if cards[0] < CARD_LEFT_MIN:
                    fails.append(
                        f"slide 9: SG-card-span leftmost {cards[0]:.2f}\" < {CARD_LEFT_MIN}\" "
                        "(Geoscan clusters cards, not edge-left)"
                    )
                if cards[-1] < CARD_RIGHT_MIN:
                    fails.append(
                        f"slide 9: SG-card-span rightmost {cards[-1]:.2f}\" < {CARD_RIGHT_MIN}\""
                    )


def check_density_ref(prs: Presentation, fails: list[str]) -> None:
    """Bodies must not be sparse stubs in Geoscan-sized slots (SG-density-ref)."""
    for (human_n, key), floor in DENSITY_FLOORS.items():
        slide = prs.slides[human_n - 1]
        body = find_shape_with(slide, lambda t, k=key: k in t)
        if not body:
            fails.append(f"slide {human_n}: SG-density-ref missing body key «{key}»")
            continue
        n = len(body.text_frame.text.strip())
        if n < floor:
            fails.append(
                f"slide {human_n}: SG-density-ref «{key}» len={n} < {floor} "
                f"(пустыня vs Geoscan-scale; удлини смысл или сверь эталон)"
            )


def check_asset_matte(fails: list[str]) -> None:
    """Portraits must not look like white-punched mattes (SG-asset-matte)."""
    try:
        from PIL import Image
    except ImportError:
        return
    for name in ("portrait_danila.png", "portrait_nikita.png", "portrait_timofey.png"):
        path = ROOT / "assets" / name
        if not path.exists():
            continue
        im = Image.open(path).convert("RGB")
        w, h = im.size
        # Center disk sample — punched mattes dump near-white into dark hair/bg
        cx, cy, r = w // 2, h // 2, min(w, h) // 3
        white = total = 0
        px = im.load()
        for y in range(cy - r, cy + r, max(1, r // 24)):
            for x in range(cx - r, cx + r, max(1, r // 24)):
                if (x - cx) ** 2 + (y - cy) ** 2 > r * r:
                    continue
                r_, g, b = px[x, y]
                total += 1
                if r_ > 245 and g > 245 and b > 245:
                    white += 1
        if total and white / total > 0.22:
            fails.append(
                f"assets/{name}: SG-asset-matte white_frac={white/total:.2f} in face disk "
                f"— re-copy Geoscan portrait verbatim (no brightness matte)"
            )


def main() -> int:
    if not PPTX.exists():
        print(f"FAIL missing {PPTX}")
        return 1

    fails: list[str] = []
    check_layout_and_assets(fails)
    check_asset_matte(fails)
    check_slide8_face_rings_no_overlap(fails)
    prs = Presentation(str(PPTX))
    check_empty_titles(prs, fails)
    check_label_contrast(prs, fails)
    check_slide8_photo_clearance(prs, fails)
    check_slide9_photo_square(prs, fails)
    check_banal_copy(prs, fails)
    check_design_gates(prs, fails)
    check_density_ref(prs, fails)

    if fails:
        print("GATE FAIL")
        for f in fails:
            print(" -", f)
        return 1

    print("GATE PASS")
    print(
        f" checked slides {list(OWN_SLIDE_INDEXES)} + "
        f"overlap/density-ref/type/overflow/cards/matte in {PPTX.name}"
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
