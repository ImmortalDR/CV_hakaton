#!/usr/bin/env python3
"""Build jury SUBMISSION: only strong content slides (junk DELETED).

Keep from draft (1-based):
  7–14  mandatory + ядро/стенд/MVP (Geoscan-quality + expert)
  16    рынок (4 колонки — читается)
  25    сквозной путь (5 шагов — читается)

Drop sparse/broken template fills: 15, 17–20, 26 and all instruction/icon slides.
"""

from __future__ import annotations

import shutil
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
SRC = ROOT / "FSP_NIKITA_DRAFT.pptx"
OUT = ROOT / "FSP_NIKITA_SUBMISSION.pptx"
KEEP = {7, 8, 9, 10, 11, 12, 13, 14, 16, 25}


def delete_slide(prs: Presentation, index: int) -> None:
    sld_id = prs.slides._sldIdLst[index]
    rId = sld_id.get(qn("r:id"))
    prs.part.drop_rel(rId)
    del prs.slides._sldIdLst[index]


def main() -> None:
    if not SRC.exists():
        raise SystemExit(f"missing {SRC}")
    shutil.copy2(SRC, OUT)
    prs = Presentation(str(OUT))
    for i in range(len(prs.slides) - 1, -1, -1):
        if (i + 1) not in KEEP:
            delete_slide(prs, i)

    # Renumber footers 1..N (draft left 7/8/…/25 — stupid jury-facing bug)
    from pptx.dml.color import RGBColor
    from pptx.util import Pt

    white = RGBColor(0xFF, 0xFF, 0xFF)
    for idx, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            name = getattr(sh, "name", "") or ""
            if "Номер слайда" not in name or not sh.has_text_frame:
                continue
            tf = sh.text_frame
            tf.clear()
            run = tf.paragraphs[0].add_run()
            run.text = str(idx)
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = white
        # Strip organizer U+000B soft-breaks (title case)
        for sh in slide.shapes:
            if not sh.has_text_frame or "\x0b" not in sh.text_frame.text:
                continue
            raw = sh.text_frame.text.replace("\x0b", "\n")
            pt, bold, rgb = 20.0, False, white
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.size is not None:
                        # snap to Geoscan ramp (organizer title often 18)
                        raw_pt = float(r.font.size.pt)
                        ramp = (11, 12, 13, 15, 16, 20, 22)
                        pt = min(ramp, key=lambda x: abs(x - raw_pt))
                    if r.font.bold is not None:
                        bold = bool(r.font.bold)
                    try:
                        if r.font.color is not None and r.font.color.type is not None:
                            rgb = r.font.color.rgb
                    except Exception:
                        pass
                    break
                break
            tf = sh.text_frame
            tf.clear()
            for li, line in enumerate(raw.split("\n")):
                p = tf.paragraphs[0] if li == 0 else tf.add_paragraph()
                run = p.add_run()
                run.text = line
                run.font.size = Pt(pt)
                run.font.bold = bold
                run.font.color.rgb = rgb

    prs.save(str(OUT))
    print(f"submission: {OUT} ({len(prs.slides)} slides)")
    print("  keep draft#:", sorted(KEEP))


if __name__ == "__main__":
    main()
