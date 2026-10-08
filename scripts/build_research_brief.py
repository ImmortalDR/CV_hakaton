"""Generate the versioned research handout; no original presentation is overwritten."""
import json
from pathlib import Path
from statistics import mean
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Pt
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

root=Path(__file__).resolve().parents[1]
r=json.loads((root/'evaluation/bank_v2_results.json').read_text())
texts=[
('Исследовательская часть ФСП · банк 2.0.0','9 октября 2026. Подготовлено по поручению Тимофея Розова. Автоматическая проверка; не заключение внешних экспертов.'),
('Что изменилось','24 семейства вместо 12; 8 вопросов вместо 4. Для подтверждения нужно 6/8 и минимум 3/4 по основному навыку Python/SQL. Старые результаты и правила старых попыток сохранены.'),
('Покрытие','Python: фильтрация, индексы, границы, дефекты, ссылки, аргументы по умолчанию, кеш, зависимости задач и правила API. Data: NULL, фильтры, соединения, группировки, окна, статистические доли и конкурентные обновления.'),
('Как измеряли','На 200 dev-вариантах каждой категории выбран частый ответ семейства. Другие 500 probe-вариантов использованы для проверки. Порог задан заранее. 42 000 симулированных попыток; 36 000 сверок с отдельными решателями, ошибок 0.'),
]
for strategy,title in [('frequent_guess','Угадывание частых ответов'),('noise_10','Случайные ошибки в 10% ответов'),('noise_20','Случайные ошибки в 20% ответов')]:
    vals=[mean(row['pass_rate'] for row in r['rows'] if row['version']==v and row['strategy']==strategy) for v in ['1.2.0','2.0.0']]
    texts.append((title,f'Средняя доля успешных попыток: {vals[0]:.2%} в 1.2.0 → {vals[1]:.2%} в 2.0.0. Это синтетическая модель, не результаты людей.'))
texts.extend([
('Что результат не доказывает','0/3000 успехов конкретного угадывания не означают защиту от списывания. Методы решения узнаваемы. Верхняя граница Wilson при 0/500 — около 0,76% в этой модели для категории. Нет доказанной эквивалентности сложности и профессиональной валидности Junior/Middle/Senior.'),
('Что готово для пилота','12 форм A/B, пустые шаблоны независимой разметки, отдельные ключи и защищённый от случайного изменения манифест. Инструмент считает согласие, матрицу ошибок, повторяемость и различия групп. Сейчас: not_run, 0 участников; неизвестные показатели остаются null.'),
('Следующий шаг команды','Содержательное ревью 24 семейств, независимые метки до тестирования, две формы без промежуточной выдачи эталонов, разбор ошибок и новая отдельная финальная выборка при изменении правил.'),
('Формулировка для выступления','Мы расширили проверяемый банк и измерили последствия на фиксированных синтетических стратегиях. Новый вариант реже пропускает частое угадывание и лучше переносит случайные ошибки в этой модели. Профессиональную валидность предстоит подтвердить пилотом, для которого подготовлены формы и расчёт метрик.'),
])
out=root/'docs/delivery';out.mkdir(exist_ok=True)
doc=Document();doc.styles['Normal'].font.name='DejaVu Sans';doc.styles['Normal'].font.size=Pt(10)
for i,(heading,body) in enumerate(texts):
    doc.add_heading(heading,level=0 if i==0 else 1);doc.add_paragraph(body)
doc.save(out/'RESEARCH_V2.docx')
pdfmetrics.registerFont(TTFont('DejaVu','/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'))
styles=getSampleStyleSheet()
for style in styles.byName.values(): style.fontName='DejaVu'
styles['BodyText'].leading=15
story=[]
for i,(heading,body) in enumerate(texts):
    story += [Paragraph(escape(heading),styles['Title'] if i==0 else styles['Heading2']),Paragraph(escape(body),styles['BodyText']),Spacer(1,9)]
SimpleDocTemplate(str(out/'RESEARCH_V2.pdf'),leftMargin=44,rightMargin=44,topMargin=40,bottomMargin=40).build(story)
print('Created RESEARCH_V2.docx and RESEARCH_V2.pdf')
