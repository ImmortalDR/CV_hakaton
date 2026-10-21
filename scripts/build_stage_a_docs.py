"""Build current Stage A documents from Markdown; preserve historical slides."""
from pathlib import Path
from xml.sax.saxutils import escape
import re
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle, PageBreak
from reportlab.lib import colors

root = Path(__file__).resolve().parents[1]
out = root / 'docs/delivery'
pdfmetrics.registerFont(TTFont('StageDejaVu', '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'))
styles = getSampleStyleSheet()
for style in styles.byName.values():
    style.fontName = 'StageDejaVu'
    style.fontSize = 9
    style.leading = 13
for name,size in [('Heading1',16),('Heading2',13),('Heading3',11)]:
    styles[name].fontSize = size
    styles[name].leading = size + 4
cell = ParagraphStyle('StageCell', parent=styles['Normal'], fontSize=8, leading=11)
story = []
doc = Document()
doc.styles['Normal'].font.name = 'DejaVu Sans'
doc.styles['Normal'].font.size = Pt(10)
def plain(value):
    return re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1 (\2)', value).replace('**', '').replace('`', '')
for index,name in enumerate(['CHANGELOG_STAGE_A.md', 'FOR_JURY.md', 'STAGE_A_SMOKE.md']):
    if index:
        story.append(PageBreak())
    source = (out / name).read_text()
    if name == 'CHANGELOG_STAGE_A.md':
        source = source.split('\n---\n', 1)[0] # old attempts remain in the source, not current reviewer brief
    lines = source.splitlines();i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                row = [plain(c.strip()) for c in lines[i].strip('|').split('|')]
                if not all(re.fullmatch(r'[:\- ]+', c) for c in row):rows.append(row)
                i += 1
            table = doc.add_table(rows=1, cols=len(rows[0]));table.style = 'Light Shading Accent 1'
            for c,value in zip(table.rows[0].cells,rows[0]):c.text = value
            for row in rows[1:]:
                for c,value in zip(table.add_row().cells,row):c.text = value
            table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            for row in table.rows:row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            pdf_table = LongTable([[Paragraph(escape(c),cell) for c in row] for row in rows], colWidths=[507/len(rows[0])]*len(rows[0]),repeatRows=1)
            pdf_table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#eee5f5')),('VALIGN',(0,0),(-1,-1),'TOP'),('GRID',(0,0),(-1,-1),.3,colors.HexColor('#ccc5d6')),('BOTTOMPADDING',(0,0),(-1,-1),6)]))
            story += [pdf_table,Spacer(1,8)];continue
        if line.startswith('#'):
            text = plain(line.lstrip('# '));level = min(3,len(line)-len(line.lstrip('#')))
            heading = doc.add_heading(text, level=level)
            if index and level == 1:
                heading.paragraph_format.page_break_before = True
            story.append(Paragraph(escape(text),styles['Heading'+str(level)]))
        elif line.strip():
            text = plain(line);doc.add_paragraph(text);story.append(Paragraph(escape(text),styles['Normal']))
        else:story.append(Spacer(1,6))
        i += 1
doc.save(out / 'STAGE_A.docx')
SimpleDocTemplate(str(out / 'STAGE_A.pdf'),pagesize=(595,842),leftMargin=44,rightMargin=44,topMargin=40,bottomMargin=40,title='ФСП — актуальная сдача этапа A',author='Нанотехнологический альянс').build(story)
print('STAGE_A DOCX/PDF generated; presentation unchanged')
