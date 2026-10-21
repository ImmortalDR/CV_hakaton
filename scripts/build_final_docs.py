"""Generate the detailed final DOCX/PDF from the current audited Markdown."""
import re
from pathlib import Path
from xml.sax.saxutils import escape
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle, PageBreak, Image
from PIL import Image as PILImage

out=Path('docs/delivery')
files=['FINAL_AUDIT_2026-10-21.md','MATCHING_CHALLENGE_2026-10-21.md','ASSESSMENT_CONTENT_REVIEW_2026-10-21.md','ACCESS_AND_OPERATIONS_AUDIT_2026-10-21.md','REMAINING_AFTER_FINAL_AUDIT_2026-10-21.md','FINAL_DEMO_AND_PITCH_2026-10-21.md']
pdfmetrics.registerFont(TTFont('FinalDejaVu','/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'))
pdfmetrics.registerFont(TTFont('FinalDejaVuBold','/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'))
styles=getSampleStyleSheet()
for name,style in styles.byName.items():
 style.fontName='FinalDejaVu';style.fontSize=9.5;style.leading=13
for name,size in [('Heading1',16),('Heading2',13),('Heading3',11)]:
 styles[name].fontName='FinalDejaVuBold';styles[name].fontSize=size;styles[name].leading=size+5
cell=ParagraphStyle('Cell',parent=styles['Normal'],fontSize=8,leading=11)
doc=Document()
doc.styles['Normal'].font.name='DejaVu Sans';doc.styles['Normal'].font.size=Pt(10)
story=[]
def plain(text):
 return re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'\1 (\2)',text).replace('**','').replace('`','')
def para(text,style='Normal'):
 return Paragraph(escape(plain(text)), styles[style])
for index,name in enumerate(files):
 if index: doc.add_page_break();story.append(PageBreak())
 lines=(out/name).read_text().splitlines();i=0
 while i<len(lines):
  line=lines[i]
  if line.startswith('|'):
   rows=[]
   while i<len(lines) and lines[i].startswith('|'):
    row=[plain(v.strip()) for v in lines[i].strip('|').split('|')]
    if not all(re.fullmatch(r'[:\- ]+',v) for v in row): rows.append(row)
    i+=1
   count=max(map(len,rows));rows=[row+['']*(count-len(row)) for row in rows]
   table=doc.add_table(rows=1,cols=count);table.style='Light Shading Accent 1'
   for j,value in enumerate(rows[0]):table.rows[0].cells[j].text=value
   for row in rows[1:]:
    for c,value in zip(table.add_row().cells,row):c.text=value
   for row in table.rows:
    for c in row.cells:
     for paragraph in c.paragraphs:
      for run in paragraph.runs: run.font.size=Pt(9)
   widths=[507/count]*count
   if count==3:widths=[96,260,151]
   if count==4:widths=[110,110,130,157]
   if count==6:widths=[95,49,77,86,91,109]
   table.autofit=False
   available=doc.sections[0].page_width-doc.sections[0].left_margin-doc.sections[0].right_margin
   for row in table.rows:
    row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    for j,c in enumerate(row.cells): c.width=int(available*widths[j]/sum(widths))
   for j,column in enumerate(table.columns): column.width=int(available*widths[j]/sum(widths))
   table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
   tab=LongTable([[Paragraph(escape(v),cell) for v in row] for row in rows],colWidths=widths,repeatRows=1,hAlign='LEFT')
   tab.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#ece7f6')),('VALIGN',(0,0),(-1,-1),'TOP'),('GRID',(0,0),(-1,-1),.3,colors.HexColor('#ccc5d6')),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),6)]))
   story += [tab,Spacer(1,9)];continue
  if line.startswith('#'):
   level=min(3,len(line)-len(line.lstrip('#')));text=line.lstrip('# ')
   doc.add_heading(plain(text),level=level);story.append(para(text,'Heading'+str(level)))
  elif line.strip():doc.add_paragraph(plain(line));story.append(para(line))
  else:story.append(Spacer(1,5))
  i+=1
story.append(PageBreak());doc.add_page_break();doc.add_heading('Проверенные экраны',level=1);story.append(para('Проверенные экраны','Heading1'))
for image_index,(name,label) in enumerate([('invitation-salary.png','Предпросмотр компании и зарплаты'),('invitation-mobile.png','Приглашение на ширине 360 px'),('mobile-390.png','Мобильное меню на ширине 390 px')]):
 if image_index: story.append(PageBreak());doc.add_page_break()
 path=out/'evidence/final_2026-10-21'/name
 doc.add_heading(label,level=2);story.append(para(label,'Heading2'))
 w,h=PILImage.open(path).size;scale=min(507/w,530/h)
 story.append(Image(str(path),width=w*scale,height=h*scale))
 doc.add_picture(str(path),width=Inches(min(6.2,6.2*w/h)))
 story.append(Spacer(1,12))
def footer(canvas,document):
 canvas.setFont('FinalDejaVu',8);canvas.drawString(44,23,'ФСП · финальный аудит · 21.10.2026');canvas.drawRightString(551,23,str(document.page))
doc.save(out/'FINAL_REPORT_2026-10-21.docx')
SimpleDocTemplate(str(out/'FINAL_REPORT_2026-10-21.pdf'),pagesize=(595,842),leftMargin=44,rightMargin=44,topMargin=40,bottomMargin=40,title='ФСП — финальный аудит 21 октября 2026',author='НТ Альянс').build(story,onFirstPage=footer,onLaterPages=footer)
print('Final DOCX/PDF generated from',len(files),'source documents')
