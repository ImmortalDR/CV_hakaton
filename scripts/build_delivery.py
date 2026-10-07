"""Create native template-based slides and standalone DOCX/PDF documentation.

The supplied template is read only. All output belongs to demo/docs/delivery.
Run from the repository root with requirements-dev.lock installed.
"""

import json
import re
import textwrap
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches as DInches, Pt as DPt, RGBColor as DColor
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    PageBreak,
)

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/delivery"
DEMO = ROOT / "demo"
PURPLE = "520978"
PINK = "FF0053"
INK = "1C1D22"
WHITE = "FFFFFF"


def slide_text(shape, text, size=20, color=INK, bold=False):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = "DejaVu Sans"
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(12)
        p.space_before = Pt(0)
    return shape


def textbox(slide, x, y, w, h, text, size=20, color=INK, bold=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    return slide_text(shape, text, size, color, bold)


def fill(
    slide, index, text, size=20, color=INK, bold=False, x=None, y=None, w=None, h=None
):
    s = slide.shapes[index]
    for key, value in [("left", x), ("top", y), ("width", w), ("height", h)]:
        if value is not None:
            setattr(s, key, Inches(value))
    return slide_text(s, text, size, color, bold)


def presentation():
    template = ROOT / "ЛЦТ_2026_Шаблон презентации ФСП(1).pptx"
    p = Presentation(template)
    s = p.slides[6]
    fill(
        s,
        0,
        "ФСП · карьера\nНавыки говорят за вас",
        34,
        WHITE,
        True,
        x=0.4,
        y=3.6,
        w=9,
        h=1.9,
    )
    fill(
        s,
        1,
        "Обратный найм: проверка навыков → категория → предложение\nMVP · две специализации · личное согласие на контакт",
        17,
        WHITE,
        x=0.4,
        y=6.05,
        w=10.5,
        h=0.95,
    )
    textbox(s, 0.5, 5.55, 8, 0.4, "Данила · Никита · Тимофей", 18, WHITE)

    s = p.slides[7]
    textbox(
        s, 0.4, 1.5, 6.1, 1.3, "Проверяемый путь\nот навыков к работе", 31, PURPLE, True
    )
    fill(
        s,
        4,
        "Данила — тимлид и разработка\nНикита — дизайн и разработка\nТимофей — исследователь\nСостав и роли: docs/TEAM_TASKS.md",
        17,
        PURPLE,
        w=6.0,
        h=2.65,
    )
    fill(s, 6, "Суть решения", 28, PURPLE, True)
    fill(
        s,
        11,
        "Кандидат проходит тест категории. Работодатель видит доказательства и приглашает лично — сразу с зарплатой.",
        20,
        INK,
        y=3.0,
        h=1.65,
    )
    fill(
        s,
        3,
        "Подтверждения, unknown и самодекларация разделены. Контакт открывается только выбранной компании.",
        19,
        INK,
        y=5.6,
        h=1.2,
    )

    # Native template layout 12: narrative left and a real application image right.
    s = p.slides[11]
    fill(s, 3, "ОТ ПРОФИЛЯ К ПРЕДЛОЖЕНИЮ", 23, PURPLE, True, w=6.2)
    fill(
        s,
        1,
        "01  Профиль и согласие\n02  Тест и подтверждённая категория\n03  Потребность и доказательства\n04  Приглашение с вилкой RUB\n05  Принятие и адресный контакт",
        23,
        PURPLE,
        y=1.65,
        h=4.7,
    )
    textbox(s, 7.35, 1.25, 5.45, 0.65, "Два кабинета · один сценарий", 20, WHITE, True)
    img = DEMO / "screenshots/06-search.png"
    s.shapes.add_picture(str(img), Inches(7.3), Inches(2.02), width=Inches(5.6))
    textbox(
        s,
        7.35,
        6.2,
        5.3,
        0.8,
        "Сохранённая подборка и источник каждого основания",
        17,
        WHITE,
    )

    # Native wide content layout 13, retaining background, logos and footer.
    s = p.slides[12]
    fill(s, 3, "ЧТО ПОДТВЕРЖДАЕТ ТЕСТ", 25, WHITE, True, w=8.8)
    fill(
        s,
        1,
        "12 семейств · 6 категорий · 4 задания на попытку\n\nСервер хранит seed, версию, эталон и результат.\nПроходной балл 75%. Неудача не понижает грейд.\nПересдача — через 24 часа; смена уровня — через 90 дней.\n\nПараметризация уменьшает пользу копирования ответов.\nЗащита от внешней помощи и валидность на людях не доказаны.",
        23,
        PURPLE,
        y=1.8,
        h=4.9,
    )

    s = p.slides[13]
    fill(s, 2, "ПРОВЕРКИ И РЕЗУЛЬТАТЫ", 23, PURPLE, True, w=5.5)
    fill(
        s,
        4,
        "36 тестов пройдено\nПрава, грейды, повторы, PDF\n\n30 пользователей · 270 чтений\n0 ошибок · p95 3,04 с\n\nDocker, новая БД, перезапуск\nи реальный Chromium-сценарий",
        23,
        PURPLE,
        y=1.45,
        h=5.2,
    )
    fill(s, 3, "", 20, INK)
    textbox(s, 7.05, 1.35, 5.7, 0.7, "P@5 на синтетическом test", 25, WHITE, True)
    textbox(s, 7.05, 2.25, 5.7, 1.1, "0,80  /  0,475", 38, WHITE, True)
    textbox(
        s,
        7.05,
        3.25,
        5.7,
        1.0,
        "MVP                       baseline\n8 запросов · 48 профилей",
        18,
        WHITE,
    )
    textbox(
        s,
        7.05,
        4.55,
        5.7,
        2.0,
        "97,92% — согласие грейдов\nна 240 синтетических классификациях.\nВнешних экспертов: 0. Пилот на людях предстоит.",
        20,
        WHITE,
    )

    # Required team slide: retain the template's native cards, resize to three people.
    s = p.slides[8]
    cards = [
        (
            12,
            15,
            22,
            21,
            "Данила",
            "Тимлид · разработка\nСервер и интеграция\nДанные и приёмка",
        ),
        (
            0,
            16,
            2,
            1,
            "Никита",
            "Дизайн · разработка\nИнтерфейс и UX\nПрезентация и демо",
        ),
        (3, 17, 5, 4, "Тимофей", "Исследования\nМетодики и эталоны\nОценка качества"),
    ]
    remove = [6, 7, 8, 9, 10, 11, 15, 16, 17, 18, 19]
    shapes = list(s.shapes)
    for k, (bg, photo, name, body, title, description) in enumerate(cards):
        x = 0.38 + k * 4.22
        shapes[bg].left = Inches(x)
        shapes[bg].width = Inches(3.96)
        for index, top, height in [
            (photo, 2.04, 1.45),
            (name, 3.78, 0.7),
            (body, 4.6, 1.9),
        ]:
            shapes[index].left = Inches(x + 0.27)
            shapes[index].width = Inches(3.42)
            shapes[index].top = Inches(top)
            shapes[index].height = Inches(height)
        textbox(s, x + 0.3, 2.1, 3.3, 1.3, f"0{k+1}", 52, PINK, True)
        slide_text(shapes[name], title, 27, PURPLE, True)
        slide_text(shapes[body], description, 20, PURPLE)
    slide_text(shapes[20], "КОМАНДА", 24, WHITE, True)
    for i in sorted(remove, reverse=True):
        e = shapes[i]._element
        e.getparent().remove(e)
    textbox(
        s,
        0.5,
        7.04,
        11.7,
        0.35,
        "Роли по TEAM_TASKS.md. Фамилии, контакты, город и организация пока не предоставлены.",
        10,
        WHITE,
    )

    s = p.slides[9]
    fill(s, 2, "ПУТЬ К MVP", 24, WHITE, True)
    fill(
        s,
        3,
        "История знакомства, организация и прошлые проекты команды не предоставлены. Здесь не заявляются вымышленные достижения.",
        20,
        PURPLE,
        w=11.3,
        h=1.02,
    )
    fill(
        s,
        7,
        "Кейс связывает проверяемые навыки с понятными условиями найма. Ценность для кандидата — контроль контактов и отсутствие потока откликов.",
        20,
        PURPLE,
        w=11.3,
        h=1.05,
    )
    fill(
        s,
        5,
        "Исправлены постоянный ответ одного семейства, историческое раскрытие стажа и нехватка памяти сборки. Проверены отрицательные сценарии и повторные запросы.",
        20,
        PURPLE,
        w=11.3,
        h=1.08,
    )

    s = p.slides[10]
    fill(s, 6, "РЕАЛИЗАЦИЯ И РАЗВИТИЕ", 23, WHITE, True)
    fill(s, 7, "Техническая суть", 24, PURPLE, True)
    fill(s, 8, "Следующий этап", 23, PURPLE, True)
    fill(
        s,
        4,
        "React / TypeScript\nFastAPI / PostgreSQL\n\nВерсионный банк и доказательства\nПодбор внутри категории\nТранзакционное согласие\nPDF по правам получателя\n\nБез платных LLM и GPU",
        21,
        PURPLE,
        h=4.5,
    )
    fill(
        s,
        5,
        "Пилот с экспертами и людьми\nРасширение банка и рубрики\nРеальные SMTP и ФСП / OIDC\nПагинация и ускорение банка\nПроцедуры работы с данными\n\nСейчас: локальный стенд и видео.\nПубличный домен не настроен.",
        21,
        PURPLE,
        h=4.5,
    )

    chosen = [6, 7, 11, 12, 13, 8, 9, 10]
    ids = list(p.slides._sldIdLst)
    for i, node in enumerate(ids):
        if i not in chosen:
            p.part.drop_rel(node.rId)
        p.slides._sldIdLst.remove(node)
    for i in chosen:
        p.slides._sldIdLst.append(ids[i])
    for number, slide in enumerate(p.slides, 1):
        for shape in slide.shapes:
            if (
                shape.has_text_frame
                and shape.text.strip().isdigit()
                and shape.top > Inches(6.8)
            ):
                old = shape.text_frame.paragraphs[0].font.color
                color = PURPLE if number in (2, 7) else WHITE
                slide_text(shape, str(number), 12, color)
        slide.notes_slide.notes_text_frame.text = (
            "MVP 1.0.0. Фактические результаты: docs/delivery/VALIDATION.md; evaluation/results.json. "
            "Имена и роли: docs/TEAM_TASKS.md. Реальный SMTP/ФСП, человеческая валидация и публичный HTTPS не подтверждены. "
            "Шаблон организаторов сохранён в исходном файле без изменений."
        )
    # Drop unused native layouts/masters so their image assets are not shipped.
    used_layouts = {slide.slide_layout.part for slide in p.slides}
    used_masters = {slide.slide_layout.slide_master.part for slide in p.slides}
    for master in list(p.slide_masters):
        if master.part in used_masters:
            for node in list(master._element.sldLayoutIdLst):
                if master.part.related_part(node.rId) not in used_layouts:
                    master.part.drop_rel(node.rId)
                    master._element.sldLayoutIdLst.remove(node)
    for node in list(p._element.sldMasterIdLst):
        if p.part.related_part(node.rId) not in used_masters:
            p.part.drop_rel(node.rId)
            p._element.sldMasterIdLst.remove(node)
    p.save(DEMO / "FSP_MVP.pptx")
    return len(p.slides)


def plain(text):
    text = re.sub(r"\[([^]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    return text.replace("**", "").replace("`", "")


def documentation():
    pdfmetrics.registerFont(
        TTFont("DV", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
    )
    pdfmetrics.registerFont(
        TTFont("DVB", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")
    )
    styles = {
        "body": ParagraphStyle(
            "body", fontName="DV", fontSize=9, leading=13, spaceAfter=7
        ),
        "h1": ParagraphStyle(
            "h1",
            fontName="DVB",
            fontSize=21,
            leading=28,
            textColor=colors.HexColor("#520978"),
            spaceAfter=16,
        ),
        "h2": ParagraphStyle(
            "h2",
            fontName="DVB",
            fontSize=13,
            leading=18,
            spaceBefore=12,
            spaceAfter=9,
            textColor=colors.HexColor("#520978"),
            keepWithNext=True,
        ),
        "cell": ParagraphStyle("cell", fontName="DV", fontSize=7, leading=10),
        "code": ParagraphStyle(
            "code",
            fontName="DV",
            fontSize=7.5,
            leading=11,
            backColor=colors.HexColor("#F4EFF8"),
            borderPadding=8,
            spaceAfter=10,
        ),
        "cover": ParagraphStyle(
            "cover",
            fontName="DVB",
            fontSize=34,
            leading=44,
            textColor=colors.HexColor("#520978"),
            spaceAfter=20,
        ),
    }
    doc = Document()
    section = doc.sections[0]
    section.top_margin = DInches(0.75)
    section.bottom_margin = DInches(0.75)
    section.left_margin = DInches(0.7)
    section.right_margin = DInches(0.7)
    normal = doc.styles["Normal"]
    normal.font.name = "DejaVu Sans"
    normal.font.size = DPt(9)
    normal.paragraph_format.space_after = DPt(7)
    for name, size in [("Title", 32), ("Heading 1", 21), ("Heading 2", 14)]:
        style = doc.styles[name]
        style.font.name = "DejaVu Sans"
        style.font.size = DPt(size)
        style.font.color.rgb = DColor.from_string(PURPLE)
    doc.add_heading("ФСП · карьера", 0)
    doc.add_paragraph("Техническая документация и отчёт проверки MVP 1.0.0")
    doc.add_paragraph("Обратный найм по подтверждённым навыкам\n7 октября 2026 · UTC")
    doc.add_paragraph(
        "Код: github.com/ImmortalDR/CV_hakaton\nРабочая ветка: codex/fsp-mvp"
    )
    doc.add_paragraph(
        "Демонстрационная реализация. Результаты синтетических измерений не являются экспертной валидацией профессиональных грейдов."
    )
    story = [
        Spacer(1, 80),
        Paragraph("ФСП · карьера", styles["cover"]),
        Paragraph(
            "Техническая документация<br/>и отчёт проверки MVP 1.0.0", styles["h1"]
        ),
        Spacer(1, 25),
        Paragraph(
            "Обратный найм по подтверждённым навыкам<br/>7 октября 2026 · UTC",
            styles["body"],
        ),
        Paragraph(
            "Код: github.com/ImmortalDR/CV_hakaton<br/>Рабочая ветка: codex/fsp-mvp",
            styles["body"],
        ),
        Spacer(1, 40),
        Paragraph(
            "Демонстрационная реализация. Результаты синтетических измерений не являются экспертной валидацией профессиональных грейдов.",
            styles["body"],
        ),
    ]
    files = [
        "ARCHITECTURE.md",
        "ASSESSMENT.md",
        "VALIDATION.md",
        "REQUIREMENTS_STATUS.md",
        "API.md",
        "FSP_INTEGRATION.md",
        "OPERATIONS.md",
    ]
    for filename in files:
        doc.add_page_break()
        story.append(PageBreak())
        lines = (OUT / filename).read_text().splitlines()
        i = 0
        while i < len(lines):
            line = lines[i]
            i += 1
            if not line.strip():
                continue
            if line.startswith("```"):
                kind = line[3:]
                block = []
                while i < len(lines) and not lines[i].startswith("```"):
                    block.append(lines[i])
                    i += 1
                i += 1
                if kind == "mermaid":
                    block = [
                        "Кандидат / работодатель → React + Nginx → FastAPI → PostgreSQL",
                        "FastAPI: банк тестов, оценка, подбор, PDF, демонстрационный провайдер ФСП.",
                    ]
                text = "\n".join(block)
                para = doc.add_paragraph(text)
                para.style = "No Spacing"
                for run in para.runs:
                    run.font.name = "DejaVu Sans Mono"
                    run.font.size = DPt(8)
                wrapped = "\n".join(
                    "\n".join(
                        textwrap.wrap(
                            l, 100, replace_whitespace=False, drop_whitespace=False
                        )
                    )
                    for l in block
                )
                story.append(
                    Paragraph(escape(wrapped).replace("\n", "<br/>"), styles["code"])
                )
            elif line.startswith("|"):
                rows = [line]
                while i < len(lines) and lines[i].startswith("|"):
                    rows.append(lines[i])
                    i += 1
                rows = [
                    [plain(cell.strip()) for cell in r.strip().strip("|").split("|")]
                    for r in rows
                    if not re.match(r"^\|[\s:|\-]+\|$", r)
                ]
                table = doc.add_table(rows=1, cols=len(rows[0]))
                table.style = "Light Shading Accent 1"
                for j, cell in enumerate(rows[0]):
                    table.rows[0].cells[j].text = cell
                for row in rows[1:]:
                    cells = table.add_row().cells
                    for j, cell in enumerate(row):
                        cells[j].text = cell
                cells = [
                    [Paragraph(escape(v), styles["cell"]) for v in row] for row in rows
                ]
                n = len(rows[0])
                widths = [487 / n] * n
                if filename == "REQUIREMENTS_STATUS.md":
                    widths = [32, 128, 182, 145]
                elif filename == "API.md":
                    widths = [200, 287]
                t = Table(cells, colWidths=widths, repeatRows=1, hAlign="LEFT")
                t.setStyle(
                    TableStyle(
                        [
                            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEE5F5")),
                            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#D9CDDF")),
                            ("VALIGN", (0, 0), (-1, -1), "TOP"),
                            ("LEFTPADDING", (0, 0), (-1, -1), 6),
                            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                            ("TOPPADDING", (0, 0), (-1, -1), 6),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                        ]
                    )
                )
                story.extend([t, Spacer(1, 12)])
            elif line.startswith("#"):
                level = len(line) - len(line.lstrip("#"))
                title = plain(line.lstrip("#").strip())
                doc.add_heading(title, min(level, 2))
                story.append(
                    Paragraph(escape(title), styles["h1" if level == 1 else "h2"])
                )
            else:
                block = [line]
                while (
                    i < len(lines)
                    and lines[i].strip()
                    and not lines[i].startswith(("#", "|", "```", "- "))
                ):
                    block.append(lines[i])
                    i += 1
                text = plain(" ".join(block))
                doc.add_paragraph(text)
                story.append(Paragraph(escape(text), styles["body"]))
    doc.save(OUT / "FSP_MVP.docx")

    def footer(canvas, document):
        canvas.setFillColor(colors.HexColor("#8A83A1"))
        canvas.setFont("DV", 8)
        canvas.drawString(44, 25, "ФСП · MVP 1.0.0 · техническая документация")
        canvas.drawRightString(550, 25, str(document.page))

    SimpleDocTemplate(
        str(OUT / "FSP_MVP.pdf"),
        title="ФСП MVP — техническая документация",
        author="CV_hakaton",
        leftMargin=44,
        rightMargin=44,
        topMargin=42,
        bottomMargin=44,
    ).build(story, onFirstPage=footer, onLaterPages=footer)


if __name__ == "__main__":
    documentation()
    slides = presentation()
    print(
        json.dumps(
            {
                "documentation": [
                    "docs/delivery/FSP_MVP.docx",
                    "docs/delivery/FSP_MVP.pdf",
                ],
                "presentation": "demo/FSP_MVP.pptx",
                "slides": slides,
            },
            ensure_ascii=False,
        )
    )
