"""Check the current guide, independently of historical slide/release checks."""
import argparse
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path

import pymupdf
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs/delivery'
SECTIONS = [
    'USER_GUIDE.md', 'ARCHITECTURE.md', 'ASSESSMENT.md', 'DATA_AND_MODELS.md',
    'VALIDATION.md', 'REQUIREMENTS_STATUS.md', 'API.md', 'FSP_INTEGRATION.md',
    'OPERATIONS.md', 'TASKS.md', 'VALIDATION_STUDIES.md',
]


def pdf_check(path):
    with pymupdf.open(path) as pdf:
        texts = []
        for index, page in enumerate(pdf, 1):
            text = page.get_text()
            assert len(text.strip()) > 80, f'Sparse/blank page {index}: {path}'
            assert '\ufffd' not in text, f'Broken glyph: {path}, {index}'
            for x0, y0, x1, y1, *_ in page.get_text('blocks'):
                assert min(x0, y0) >= 0 and x1 <= page.rect.width + 1 and y1 <= page.rect.height + 1, (path, index, 'overflow')
            texts.append(text)
        return len(pdf), ' '.join(' '.join(texts).split())


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--docx-preview', type=Path)
    args = parser.parse_args()
    pages, pdf_text = pdf_check(OUT / 'FSP_MVP.pdf')
    doc = Document(OUT / 'FSP_MVP.docx')
    doc_text = ' '.join(' '.join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells]).split())
    for section in SECTIONS:
        title = (OUT / section).read_text().splitlines()[0].lstrip('# ')
        assert title in pdf_text, ('PDF missing section', title)
        assert title in doc_text, ('DOCX missing section', title)
    for term in ['11 октября 2026', '2026-10-11-modernization', 'matching-1.1.0', '84 passed', 'NDCG@10', '24 часа', '90 дней', 'assessment-availability', 'production_promotion_allowed=false']:
        assert term in pdf_text and term in doc_text, ('Missing current fact', term)
    for obsolete in ['публичного домена нет', 'Реальные CV для разработки не использовались', '12 семейств', 'Рабочая ветка: codex/fsp-mvp']:
        assert obsolete not in pdf_text and obsolete not in doc_text, ('Obsolete claim', obsolete)
    api = json.loads((OUT / 'openapi.json').read_text())
    operations = sum(method in {'get', 'post', 'put', 'patch', 'delete'} for path in api['paths'].values() for method in path)
    models = len(api['components']['schemas'])
    assert f'{operations} операций / {models} моделей' in (OUT / 'REQUIREMENTS_STATUS.md').read_text()
    result = {
        'status': 'pass', 'checked_at': datetime.now(timezone.utc).isoformat(),
        'release': '2026-10-11-modernization',
        'scope': 'Documentation content, PDF text bounds, DOCX structure and OpenAPI counts; not a rerun of application/model tests',
        'pdf_pages': pages, 'docx_paragraphs': len(doc.paragraphs),
        'docx_tables': len(doc.tables), 'sections': SECTIONS,
        'openapi_operations': operations, 'openapi_models': models,
        'docx_render': 'not_run',
    }
    if args.docx_preview:
        preview_pages, preview_text = pdf_check(args.docx_preview)
        for section in SECTIONS:
            title = (OUT / section).read_text().splitlines()[0].lstrip('# ')
            assert title in preview_text, ('Rendered DOCX missing section', title)
        result['docx_render'] = {'status': 'pass', 'pages': preview_pages}
    result['artifacts'] = {
        name: {'bytes': (OUT / name).stat().st_size, 'sha256': hashlib.sha256((OUT / name).read_bytes()).hexdigest()}
        for name in ['FSP_MVP.pdf', 'FSP_MVP.docx', *SECTIONS]
    }
    (OUT / 'documentation-checks.json').write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps({k: v for k, v in result.items() if k != 'artifacts'}, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
