"""Validate current documentation exports; does not rerun application tests.

Usage: .venv/bin/python scripts/check_current_docs.py --rendered-docx DIR
DIR contains PDFs converted from the three DOCX files by LibreOffice.
Historical release-checks.json and historical test results are not overwritten.
"""

import argparse
import hashlib
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import unquote

import pymupdf
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/delivery"


def inspect_pdf(path):
    with pymupdf.open(path) as pdf:
        assert len(pdf), path
        for number, page in enumerate(pdf, 1):
            assert page.get_text().strip() or page.get_images(), (path, number, "empty")
            for x0, y0, x1, y1, *_ in page.get_text("blocks"):
                assert x0 >= 0 and y0 >= 0 and x1 <= page.rect.width + 1 and y1 <= page.rect.height + 1, (path, number, "bounds")
        assert "119" in "\n".join(page.get_text() for page in pdf), path
        return len(pdf)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--rendered-docx", type=Path, required=True)
    args = parser.parse_args()
    api = json.loads((OUT / "openapi.json").read_text())
    methods = {"get", "post", "put", "patch", "delete", "head", "options", "trace"}
    operations = sum(key in methods for item in api["paths"].values() for key in item)
    schemas = len(api["components"]["schemas"])
    assert (operations, schemas) == (37, 49)
    evidence = {}
    for filename, count in [("auth-browser.json", 8), ("live-smoke.json", 24)]:
        path = OUT / "evidence/human_followup_2026-10-21" / filename
        result = json.loads(path.read_text())
        assert len(result["checks"]) == count and not result.get("page_errors")
        if filename == "live-smoke.json":
            assert result["status"] == "pass"
        evidence[str(path.relative_to(ROOT))] = {"previously_recorded_checks": count}
    backlog = json.loads((ROOT / "docs/final_backlog_2026-10-20/BACKLOG_STAGE_A.json").read_text())
    tasks = {task["id"]: task["status"] for task in backlog["tasks"]}
    assert all(tasks[f"Q{i}"] == "done" for i in range(1, 8))
    assert tasks["Pck2"] == "partial" and tasks["Pck5"] != "done"
    artifacts = {}
    for name in ["FSP_MVP", "FINAL_REPORT_2026-10-21", "STAGE_A"]:
        doc = Document(OUT / f"{name}.docx")
        assert doc.paragraphs and doc.tables, name
        for extension in ["pdf", "docx"]:
            path = OUT / f"{name}.{extension}"
            artifacts[str(path.relative_to(ROOT))] = {
                "bytes": path.stat().st_size,
                "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            }
        artifacts[f"docs/delivery/{name}.pdf"]["pages"] = inspect_pdf(OUT / f"{name}.pdf")
        artifacts[f"docs/delivery/{name}.docx"]["rendered_pages"] = inspect_pdf(args.rendered_docx / f"{name}.pdf")
    # Current maintained documents only: historical snapshots can refer to old bundles.
    names = ["CURRENT_STATUS_2026-10-21", "API", "ARCHITECTURE", "OPERATIONS",
             "REQUIREMENTS_STATUS", "VALIDATION", "HUMAN_ACCEPTANCE_2026-10-21",
             "HUMAN_VERIFY_PROTOCOL_STAGE_A", "REMAINING_AFTER_FINAL_AUDIT_2026-10-21",
             "FOR_JURY", "SUBMISSION_LINKS", "GIT_DELIVERY", "STAGE_A_SMOKE"]
    sources = [ROOT / "README.md", ROOT / "demo/SCENARIO.md"] + [OUT / f"{name}.md" for name in names]
    for path in sources:
        for target in re.findall(r"\[[^\]]*\]\(([^)]+)\)", path.read_text()):
            if "://" in target or target.startswith(("#", "mailto:")):
                continue
            local = unquote(target.split("#", 1)[0])
            assert (path.parent / local).exists(), (path, target)
    result = {
        "status": "pass", "checked_at": datetime.now(timezone.utc).isoformat(),
        "scope": "Documentation structure, local links, PDF bounds/nonempty pages, DOCX rendering and existing evidence consistency. No new functional test run or external publication.",
        "visual_review": "Representative pages inspected; automated bounds/nonempty checks cover all exported and DOCX-rendered pages.",
        "openapi_operations": operations, "openapi_schemas": schemas,
        "markdown_files_checked": len(sources), "existing_evidence": evidence,
        "artifacts": artifacts,
    }
    (OUT / "documentation-checks.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
